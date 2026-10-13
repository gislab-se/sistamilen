"""Import the supplied trip archive into a local, reproducible review map.

Default builds use cached local maps and address matches. An explicit flag can
refresh public SCB polygons. Trips are observations, not reconstructed GPS tracks.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import warnings
import zipfile
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from collections import Counter
from datetime import datetime, time, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook
from pyproj import Transformer

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "prototypes" / "fredriksberg_resor"
DEFAULT_SOURCE = ROOT / "data" / "fredriksberg" / "fwdpaketleveranser.zip"
EXTRA_PLACES = APP / "source" / "scb_extra_places_raw.geojson"


def refresh_extra_places():
    """Fetch public place polygons, sending no trip or address information."""
    towns = ["Fagersta", "Kungsgården", "Sala", "Storvik", "Storvreta", "Uppsala"]
    query = {"service": "WFS", "version": "1.1.0", "request": "GetFeature", "typeName": "stat:Tatorter_2023",
             "outputFormat": "application/json", "srsName": "EPSG:3006",
             "CQL_FILTER": "tatort IN (" + ",".join("'" + t + "'" for t in towns) + ")"}
    url = "https://geodata.scb.se/geoserver/stat/wfs?" + urlencode(query)
    with urlopen(Request(url, headers={"User-Agent": "Dalarna-local-review/1.0"}), timeout=40) as response:
        raw = response.read()
    payload = json.loads(raw)
    if payload.get("type") != "FeatureCollection" or not payload.get("features"):
        raise ValueError("SCB did not return public place polygons")
    EXTRA_PLACES.parent.mkdir(parents=True, exist_ok=True)
    EXTRA_PLACES.write_bytes(raw)
    metadata = {"url": url, "source": "SCB stat:Tatorter_2023", "crs": "EPSG:3006", "reference_year": 2023,
                "retrieved_at_utc": datetime.now(timezone.utc).isoformat(), "sha256": hashlib.sha256(raw).hexdigest()}
    (EXTRA_PLACES.parent / "scb_extra_places_metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")


def normalize(value):
    return " ".join(str(value or "").casefold().split())


def number(value):
    if value in (None, ""):
        return None
    return float(str(value).replace(",", "."))


def seconds(value):
    if value in (None, ""):
        return None
    if isinstance(value, timedelta):
        return round(value.total_seconds(), 3)
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    return round(float(value) * 86400, 3)


def address_parts(address):
    match = re.match(r"^(.*),\s*(\d{3})\s*(\d{2})\s+(.+)$", address)
    if match:
        return match[1].strip(), match[2] + match[3], match[4].strip()
    return address, "", ""


def ring_centroid(ring):
    pairs = list(zip(ring, ring[1:]))
    area2 = sum(a[0] * b[1] - b[0] * a[1] for a, b in pairs)
    if abs(area2) < 1e-12:
        return 0, ring[0]
    x = sum((a[0] + b[0]) * (a[0] * b[1] - b[0] * a[1]) for a, b in pairs) / (3 * area2)
    y = sum((a[1] + b[1]) * (a[0] * b[1] - b[0] * a[1]) for a, b in pairs) / (3 * area2)
    return abs(area2), [round(x, 6), round(y, 6)]


def local_places():
    source = ROOT / "data" / "derived" / "scb_platsomraden_dalarna.geojson"
    places = {}
    for feature in json.loads(source.read_text(encoding="utf-8"))["features"]:
        props = feature["properties"]
        if not props.get("namn") or props["omradestyp"] != "Tätort":
            continue
        geometry = feature["geometry"]
        polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
        area, point = max((ring_centroid(p[0]) for p in polygons), key=lambda item: item[0])
        key = normalize(props["namn"])
        if key not in places or area > places[key]["area"]:
            places[key] = {"name": props["namn"], "coordinate": point, "area": area,
                           "municipality": props["kommunnamn"], "year": props["referensar"],
                           "source": source.relative_to(ROOT).as_posix()}
    if EXTRA_PLACES.exists():
        transform = Transformer.from_crs(3006, 4326, always_xy=True)
        for feature in json.loads(EXTRA_PLACES.read_text(encoding="utf-8-sig"))["features"]:
            props = feature["properties"]
            geometry = feature["geometry"]
            polygons = geometry["coordinates"] if geometry["type"] == "MultiPolygon" else [geometry["coordinates"]]
            area, point = max((ring_centroid(p[0]) for p in polygons), key=lambda item: item[0])
            lon, lat = transform.transform(*point)
            places[normalize(props["tatort"])] = {"name": props["tatort"], "coordinate": [round(lon,6),round(lat,6)],
                                                   "area": area, "municipality": props.get("kn"), "year": 2023,
                                                   "source": "SCB stat:Tatorter_2023, source/scb_extra_places_raw.geojson"}
    if normalize("Floda / Dala-Floda") in places:
        places[normalize("Dala-Floda")] = places[normalize("Floda / Dala-Floda")]
    fredriksberg = [p for p in places.values() if p["name"].startswith("Fredriksberg ")]
    if fredriksberg:
        weight = sum(p["area"] for p in fredriksberg)
        places[normalize("Fredriksberg")] = {
            **fredriksberg[0], "name": "Fredriksberg", "area": weight,
            "coordinate": [round(sum(p["coordinate"][i] * p["area"] for p in fredriksberg) / weight, 6) for i in range(2)],
            "source": source.relative_to(ROOT).as_posix() + " (östra och västra, areaviktat ortankare)",
        }
    # A postal locality can span only part of an SCB urban area. Use existing
    # service points as an explicit locality anchor, never as a trip address.
    with (ROOT / "data" / "derived" / "servicenoder_2026.csv").open(encoding="utf-8-sig") as stream:
        groups = {}
        for row in csv.DictReader(stream):
            groups.setdefault(normalize(row["postort"]), []).append(row)
    for key, rows in groups.items():
        if key not in places:
            places[key] = {"name": rows[0]["postort"], "area": 0,
                           "coordinate": [sum(float(r[k]) for r in rows) / len(rows) for k in ("lon", "lat")],
                           "municipality": rows[0]["kommun"], "year": 2026,
                           "source": "data/derived/servicenoder_2026.csv (medelpunkt för postortens paketnoder)"}
    return places


def import_archive(source):
    trips, reports, files, flags = [], [], [], []
    with zipfile.ZipFile(source) as archive:
        for filename in sorted(archive.namelist()):
            if not filename.lower().endswith(".xlsx"):
                continue
            raw = archive.read(filename)
            with archive.open(filename) as stream:
                with warnings.catch_warnings():
                    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
                    workbook = load_workbook(stream, read_only=True, data_only=True)
                rows = list(workbook.worksheets[0].iter_rows(values_only=True))
                headers = rows[1]
                kind = "month" if len(workbook.sheetnames) == 1 else "day"
                periods = re.findall(r"\d{4}-\d{2}-\d{2}", filename)
                file_id = "_".join(periods)
                files.append({"id": file_id, "file": filename, "kind": kind,
                              "period_start": periods[0], "period_end": periods[1],
                              "sha256": hashlib.sha256(raw).hexdigest()})
                for row_num, row in enumerate(rows[2:], 3):
                    entry = dict(zip(headers, row))
                    vehicle = entry.get("Registreringsnummer")
                    if not vehicle:
                        continue
                    reports.append({"source_id": file_id, "source_row": row_num, "kind": kind,
                                    "vehicle": vehicle, "date": periods[0],
                                    "trip_count": int(number(entry.get("Antal resor")) or 0),
                                    "distance_km": number(entry.get("Körsträcka")),
                                    "co2_source": number(entry.get("CO2"))})
                for sheet in workbook.worksheets[1:]:
                    header_found = False
                    columns = {}
                    for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if row[0] == "Start":
                            columns = {str(v): i for i, v in enumerate(row) if v is not None}
                            header_found = True
                            continue
                        if not header_found or not isinstance(row[0], datetime):
                            continue
                        get = lambda name: row[columns[name]]
                        stop = get("Stopp")
                        if not isinstance(stop, datetime):
                            raise ValueError(f"Invalid stop time: {filename}, {sheet.title}, {row_num}")
                        identity = f"{sheet.title}|{row[0].isoformat()}|{stop.isoformat()}|{get('Från')}|{get('Till')}"
                        trip = {"id": hashlib.sha256(identity.encode()).hexdigest()[:16],
                                "vehicle": sheet.title, "start": row[0].isoformat(timespec="milliseconds"),
                                "end": stop.isoformat(timespec="milliseconds"), "date": row[0].date().isoformat(),
                                "origin": str(get("Från") or ""), "destination": str(get("Till") or ""),
                                "distance_km": number(get("Sträcka")), "travel_seconds": seconds(get("Restid")),
                                "stop_seconds": seconds(get("Stopptid")),
                                "source_id": file_id, "source_sheet": sheet.title, "source_row": row_num}
                        if stop < row[0]:
                            flags.append({"type": "negative_duration", "trip_id": trip["id"]})
                        trips.append(trip)
                workbook.close()
    unique = {}
    for trip in trips:
        if trip["id"] in unique:
            flags.append({"type": "duplicate_trip", "trip_id": trip["id"], "source_id": trip["source_id"]})
        else:
            unique[trip["id"]] = trip
    trips = sorted(unique.values(), key=lambda t: (t["start"], t["vehicle"]))
    counts = Counter((t["source_id"], t["vehicle"]) for t in trips)
    for report in reports:
        if report["kind"] == "day":
            report["detail_count"] = counts[(report["source_id"], report["vehicle"])]
            report["difference"] = report["trip_count"] - report["detail_count"]
            if report["difference"]:
                flags.append({"type": "summary_difference", "source_id": report["source_id"],
                              "vehicle": report["vehicle"], "reported": report["trip_count"],
                              "imported": report["detail_count"], "difference": report["difference"]})
    return trips, reports, files, flags


def ring_contains(point, ring):
    inside=False
    for a,b in zip(ring,ring[1:]):
        if (a[1]>point[1])!=(b[1]>point[1]):
            x=a[0]+(point[1]-a[1])*(b[0]-a[0])/(b[1]-a[1])
            if point[0]<x:inside=not inside
    return inside


def reserve_centroid(place, places):
    """Use the enclosing named SCB urban area for postal-locality anchors."""
    if 'medelpunkt' not in place['source']:
        return place['coordinate'],place['source']
    path=ROOT/'data'/'derived'/'scb_platsomraden_dalarna.geojson'
    for feature in json.loads(path.read_text(encoding='utf-8'))['features']:
        props=feature['properties']
        if props['omradestyp']!='Tätort':continue
        geometry=feature['geometry']
        polygons=geometry['coordinates'] if geometry['type']=='MultiPolygon' else [geometry['coordinates']]
        if any(ring_contains(place['coordinate'],p[0]) and not any(ring_contains(place['coordinate'],hole) for hole in p[1:]) for p in polygons):
            area=places.get(normalize(props['namn']))
            if area:return area['coordinate'],f"SCB tätort 2023, centroid för {props['namn']} (postort {place['name']})"
    return place['coordinate'],place['source']


def name_vehicle_profiles(trips, reports, flags):
    """Replace source registration/sheet identifiers before exporting app data.

    Names describe frequently observed localities, not a fixed itinerary or an
    asserted operator. The source-to-alias key stays outside the served app.
    """
    originals = sorted({r['vehicle'] for r in trips + reports})
    aliases = {original: f'transport-{number:02d}' for number, original in enumerate(originals, 1)}
    colors = ['#2864a0', '#b25a31', '#2e836b', '#8456a5', '#b28221', '#b14d7c', '#265d59',
              '#6568b6', '#7c7c26', '#bd4549', '#477e9f', '#89533c', '#65864b', '#996399',
              '#9b6829', '#4c6595', '#3e9185', '#a54e64', '#67713b', '#625280']
    profiles = []
    key = {}
    for number, original in enumerate(originals, 1):
        observed = [t for t in trips if t['vehicle'] == original]
        counts = Counter(p for t in observed for p in set(address_parts(a)[2] for a in (t['origin'], t['destination'])) if p)
        main = [p for p, _ in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:3]]
        if 'Fredriksberg' in counts:
            main = ['Fredriksberg'] + [p for p in main if p != 'Fredriksberg'][:2]
        label = f"{' / '.join(main) or 'Transport'} · T{number:02d}"
        profiles.append({'id': aliases[original], 'label': label, 'short_name': f'T{number:02d}',
                         'localities': main, 'color': colors[(number-1) % len(colors)],
                         'trip_count': len(observed), 'fredriksberg': 'Fredriksberg' in counts})
        key[aliases[original]] = {'label': label, 'original_vehicle_and_sheet': original}
    for row in trips + reports + flags:
        if 'vehicle' in row:
            row['vehicle'] = aliases[row['vehicle']]
        if 'source_sheet' in row:
            row['source_sheet'] = row['vehicle']
    private = ROOT / '.runtime' / 'fredriksberg-review'
    private.mkdir(parents=True, exist_ok=True)
    (private / 'vehicle-key.json').write_text(json.dumps(key, ensure_ascii=False, indent=2), encoding='utf-8')
    return profiles


def build(source, coordinates=None):
    trips, reports, files, flags = import_archive(source)
    places = local_places()
    overrides = {}
    if coordinates:
        with coordinates.open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                lon, lat = float(row["lon"]), float(row["lat"])
                if not (10 < lon < 25 and 54 < lat < 70) or not row.get("source"):
                    raise ValueError("Coordinates must be in Sweden and include a source")
                overrides[normalize(row["address"])] = {"coordinate": [lon, lat], "precision": "address",
                                                        "source": row["source"]}
    matches_path = APP / "source" / "address_matches.json"
    match_cache = json.loads(matches_path.read_text(encoding="utf-8")) if matches_path.exists() else {}
    matches = match_cache.get("addresses", {})
    public_places_path=APP/"source"/"public_place_coordinates.json"
    public_places=json.loads(public_places_path.read_text(encoding="utf-8")) if public_places_path.exists() else {}
    addresses = {}
    for address in sorted({a for t in trips for a in (t["origin"], t["destination"])}):
        street, postcode, locality = address_parts(address)
        place = places.get(normalize(locality))
        item = {"address": address, "street": street, "postcode": postcode, "locality": locality,
                "coordinate": None, "precision": "unmapped", "source": None}
        if place:
            item.update(coordinate=place["coordinate"], precision="locality", source=place["source"],
                        municipality=place["municipality"], reference_year=place["year"])
        item["locality_coordinate"] = item["coordinate"]
        item["locality_source"] = item["source"]
        item["geocoding"] = dict(matches.get(address, {"coordinate": None, "status": "unmatched", "reason": "Ingen adressmatchning tillgänglig", "source": None}))
        if address in public_places.get('addresses',{}):
            item["geocoding"]=dict(public_places['addresses'][address])
        if normalize(address) in overrides:
            item["geocoding"] = {**overrides[normalize(address)], "status": "manual", "reason": "Angiven granskningskoordinat"}
        if not item['geocoding'].get('coordinate') and item['locality_coordinate']:
            reserve_point,reserve_source=reserve_centroid(place,places)
            item['geocoding'].update(coordinate=reserve_point,precision='locality',status='fallback',
                                     reason='Ingen entydig adress-/platsträff. Punkten ligger i ortens centroid eller dokumenterade ortankare och visar inte besöksadressen',
                                     source=reserve_source,geometry='locality_centroid_fallback',
                                     unresolved_reason=item['geocoding']['reason'])
        item["address_coordinate"] = item["geocoding"].get("coordinate")
        addresses[address] = item
    nodes = []
    with (ROOT / "data" / "derived" / "servicenoder_2026.csv").open(encoding="utf-8-sig") as stream:
        for row in csv.DictReader(stream):
            nodes.append({"id": row["kluster_id"], "name": row["nodnamn"], "address": row["adress"],
                          "locality": row["postort"], "actors": row["aktorer"],
                          "coordinate": [float(row["lon"]), float(row["lat"])]})
    precision = Counter(a["precision"] for a in addresses.values())
    vehicles = name_vehicle_profiles(trips, reports, flags)
    data = {"meta": {"source": source.relative_to(ROOT).as_posix(), "archive_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                     "timezone": "Europe/Stockholm", "trip_count": len(trips), "address_count": len(addresses),
                     "precision_counts": dict(precision), "coordinate_method": "Ortankare: SCB tätort 2023, centroid för största delpolygonen; Fredriksbergs delar areaviktas. Postorter utan namngiven SCB-tätort använder medelpunkt för lokala paketnoder 2026. Adresspunkter redovisas separat.",
                     "distance_method": "Körsträcka från källfilen, inte beräknad från kartan",
                     "address_match_counts": dict(Counter(a["geocoding"]["status"] for a in addresses.values())),
                     "address_precision_counts": dict(Counter(a["geocoding"].get("precision","unmapped") for a in addresses.values())),
                     "address_match_method": "Adress-/platsträffar från lokala koordinater, OpenStreetMap, bredare Nominatim-sökning och officiella kontaktkort. Osäkra gatupunkter och verksamhetsträffar redovisas separat. Om adressen inte hittas används ortcentroid eller dokumenterat ortankare som tydligt markerad reserv.",
                     "address_match_source": match_cache.get("metadata"),
                     "co2_unit": "Ej angiven i Excel-uttaget"},
            "files": files, "trips": trips, "reports": reports, "addresses": addresses, "nodes": nodes, "flags": flags,
            "vehicles": vehicles}
    APP.mkdir(parents=True, exist_ok=True)
    out = APP / "data"
    out.mkdir(exist_ok=True)
    (out / "review-data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "review-data.js").write_text("window.REVIEW_DATA = " + json.dumps(data, ensure_ascii=False) + ";\n", encoding="utf-8")
    background = json.loads((ROOT / "data" / "derived" / "deso_2025_dalarna.geojson").read_text(encoding="utf-8"))
    for feature in background["features"]:
        feature["properties"] = {}
    (out / "geography.js").write_text("window.REVIEW_GEOGRAPHY = " + json.dumps(background, separators=(",", ":")) + ";\n", encoding="utf-8")
    match_rows=[{"address":a["address"],"status":a["geocoding"]["status"],"precision":a["geocoding"].get("precision","unmapped"),
                 "reason":a["geocoding"]["reason"],"source":a["geocoding"]["source"],
                 "lon":a["address_coordinate"][0] if a["address_coordinate"] else "",
                 "lat":a["address_coordinate"][1] if a["address_coordinate"] else ""} for a in addresses.values()]
    for name, rows in (("resor.csv", trips), ("summeringar.csv", reports), ("adresser.csv", list(addresses.values())), ("adressmatchning.csv",match_rows)):
        columns = list(dict.fromkeys(k for row in rows for k in row))
        with (out / name).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            writer.writerows(rows)
    vendor = ROOT / "prototypes" / "dva_transportnarvaro" / "vendor"
    shutil.copytree(vendor, APP / "vendor", dirs_exist_ok=True)
    print(json.dumps({"trips": len(trips), "vehicles": len({t['vehicle'] for t in trips}),
                      "addresses": len(addresses), "precision": dict(precision),
                      "reported_daily_trips": sum(r['trip_count'] for r in reports if r['kind'] == 'day'),
                      "quality_flags": len(flags)}, ensure_ascii=True))
    return data


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--coordinates", type=Path, help="Reviewed CSV: address,lon,lat,source")
    parser.add_argument("--refresh-places", action="store_true", help="Fetch six public SCB place polygons outside Dalarna")
    args = parser.parse_args()
    if args.refresh_places:
        refresh_extra_places()
    build(args.source.resolve(), args.coordinates)
